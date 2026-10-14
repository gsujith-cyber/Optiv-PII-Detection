"""pipeline.py (Member 4) - connects extractors -> detection -> redaction.

Teammates' modules are called through the HOOKS below. If a hook is missing or
its module is not ready yet, a TEMPORARY FALLBACK in this file is used so the
UI can be built and demoed today. When teammates finish, nothing changes here
except (if needed) the names in the HOOKS section.

Agreed shapes:
  extract(path)            -> [{"page": int, "text": str}, ...]
  detect(pages)            -> [{"type","value","page","start","end"}, ...]
  redact(pages, findings)  -> {"redacted_pages": [...same shape as pages...],
                               "mapping": {"[PERSON_1]": "John A. Smith", ...}}
"""
import hashlib
import importlib
import inspect
import json
import os
import re
import threading
import time
import uuid
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path

# ---------------------------------------------------------------- HOOKS ----
# (module path, function name). Change a name here if a teammate uses another.
EXTRACTOR_HOOKS = {
    ".pdf": ("extractors.pdf_extractor", "extract"),
    ".docx": ("extractors.docx_extractor", "extract"),
    ".pptx": ("extractors.pptx_extractor", "extract"),
    ".png": ("extractors", "extract_image"),
    ".jpg": ("extractors", "extract_image"),
    ".jpeg": ("extractors", "extract_image"),
}
DETECTOR_HOOKS = [   # (module, function, used_in_quick_verify)
    ("detection.regex_detector", "detect", True),    # Member 2 - regex layer (fast)
    ("processing.ner_detector", "detect", False),    # Member 2 - name/NER layer (slower)
]
REDACTOR_HOOK = ("processing.redactor", "redact")  # Member 3

SUPPORTED = {".pdf": "PDF", ".docx": "Word (DOCX)", ".pptx": "PowerPoint (PPTX)", ".txt": "Text (TXT)",
             ".png": "Image (OCR)", ".jpg": "Image (OCR)", ".jpeg": "Image (OCR)"}

# Set True if company / organisation names should also be redacted (default: people only).
REDACT_ORGANIZATIONS = False

DATA = Path("data")
ORIG_DIR, RED_DIR, CACHE_DIR = DATA / "original", DATA / "redacted", DATA / "cache"
for _d in (ORIG_DIR, RED_DIR, CACHE_DIR):
    _d.mkdir(parents=True, exist_ok=True)

OCR_DPI = 150          # lower = faster OCR, higher = more accurate (150-200 is a good range)
_progress_cb = None


def _report(frac, msg, stage=None):
    """Send progress to the UI: (0..1, text, stage name)."""
    if _progress_cb:
        try:
            _progress_cb(min(max(frac, 0.0), 1.0), msg, stage)
        except Exception:
            pass


def _hook(module, func):
    try:
        return getattr(importlib.import_module(module), func)
    except (ImportError, AttributeError):
        return None


def detect_type(filename: str):
    return SUPPORTED.get(Path(filename).suffix.lower())


# ----------------------------------------------------------- helpers -------
def _norm_pages(res):
    if isinstance(res, str):
        return [{"page": 1, "text": res}]
    out = []
    for i, p in enumerate(res, 1):
        out.append({"page": i, "text": p} if isinstance(p, str) else
                   {"page": p.get("page", i), "text": p.get("text", "")})
    return out


def _norm_findings(items):
    out = []
    for f in items or []:
        out.append({
            "type": str(f.get("type") or f.get("entity_type") or f.get("label") or "PII").upper(),
            "value": f.get("value") or f.get("text") or "",
            "page": f.get("page", 1),
            "start": f["start"], "end": f["end"],
        })
    return out


def _resolve_overlaps(findings):
    """When findings overlap on a page, keep one: specific IDs (email, SSN, PAN...)
    win over name guesses (PERSON); otherwise the longer finding wins."""
    kept = []
    findings = [f for f in findings if not _PH_ONLY.fullmatch(f["value"].strip())]
    ranked = sorted(findings, key=lambda f: (f["type"] == "PERSON", -(f["end"] - f["start"])))
    for f in ranked:
        if any(k["page"] == f["page"] and f["start"] < k["end"] and k["start"] < f["end"] for k in kept):
            continue
        kept.append(f)
    return sorted(kept, key=lambda f: (f["page"], f["start"]))


# ------------------------------------------- TEMPORARY FALLBACKS (delete later)
def _fallback_extract(path: Path):
    ext = path.suffix.lower()
    if ext == ".txt":
        return [{"page": 1, "text": path.read_text(errors="ignore")}]
    if ext == ".pdf":
        try:  # PyMuPDF is much faster than pdfplumber
            import pymupdf
            with pymupdf.open(str(path)) as doc:
                pages = [{"page": i, "text": pg.get_text("text")} for i, pg in enumerate(doc, 1)]
        except ImportError:
            import pdfplumber
            with pdfplumber.open(path) as pdf:
                pages = [{"page": i, "text": pg.extract_text() or ""} for i, pg in enumerate(pdf.pages, 1)]
        # Pages with (almost) no text are scanned images -> OCR them.
        scanned = [p["page"] for p in pages if len(p["text"].strip()) < 20]
        if scanned:
            ocr = _ocr_pdf_pages(path, scanned)
            for p in pages:
                if p["page"] in ocr:
                    p["text"] = ocr[p["page"]]
        return pages
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for r in t.rows:
                parts.append(" | ".join(c.text for c in r.cells))
        return [{"page": 1, "text": "\n".join(parts)}]
    if ext == ".pptx":
        from pptx import Presentation
        pages = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            parts = []
            for sh in slide.shapes:
                if sh.has_text_frame:
                    parts.append(sh.text_frame.text)
                if getattr(sh, "has_table", False) and sh.has_table:
                    for r in sh.table.rows:
                        parts.append(" | ".join(c.text for c in r.cells))
            pages.append({"page": i, "text": "\n".join(parts)})
        return pages
    raise ValueError("This file type needs the extractors package (extractors/__init__.py).")


def _clean_ocr(text: str) -> str:
    """Safety net for OCR line-break damage: re-join an email split as 'name@site.' / 'com'
    (only when the second part is a common domain ending)."""
    gap = r"(?:[ \t]*\n){1,2}[ \t]*"
    tld = r"(com|org|net|edu|gov|in|co|io|uk|us|ai)\b"
    text = re.sub(r"([\w.+-]+@[\w-]+(?:\.[\w-]+)*\.)" + gap + tld, r"\1\2", text)
    text = re.sub(r"([\w.+-]+@[\w-]+)" + gap + r"(\.(?:com|org|net|edu|gov|in|co|io|uk|us|ai)\b)", r"\1\2", text)
    return text


def _ocr_pdf_pages(path: Path, page_numbers):
    """OCR the given 1-based PDF pages (parallel + cached).
    Needs: pymupdf, pytesseract, pillow + the Tesseract program."""
    try:
        import pymupdf
        import pytesseract
        from PIL import Image
    except ImportError:
        raise ValueError("This looks like a scanned file. Install OCR libraries: "
                         "pip install pymupdf pytesseract pillow")
    win = r"C:\Program Files\Tesseract-OCR\tesseract.exe"
    if os.name == "nt" and os.path.exists(win):
        pytesseract.pytesseract.tesseract_cmd = win
    try:
        pytesseract.get_tesseract_version()
    except Exception:
        raise ValueError("Tesseract OCR program not found. Install it (Windows: UB-Mannheim "
                         "installer) and restart the app.")

    os.environ.setdefault("OMP_THREAD_LIMIT", "1")   # one thread per Tesseract, many in parallel
    key = hashlib.sha1(path.read_bytes()).hexdigest()
    cache_file = CACHE_DIR / f"{key}_ocr{OCR_DPI}_p4.json"
    cache = json.loads(cache_file.read_text()) if cache_file.exists() else {}
    todo = [n for n in page_numbers if str(n) not in cache]
    total, done = len(page_numbers), len(page_numbers) - len(todo)
    _report(0.02, f"OCR: {done}/{total} pages ready" + (" (from cache)" if not todo else ""), "extract")

    if todo:
        workers = min(8, os.cpu_count() or 4)

        def ocr(img):
            return pytesseract.image_to_string(img, config="--oem 1 --psm 4")

        with pymupdf.open(str(path)) as doc, ThreadPoolExecutor(workers) as pool:
            for i in range(0, len(todo), workers * 2):
                batch, imgs = todo[i:i + workers * 2], []
                for n in batch:   # render in the main thread (pymupdf is not thread-safe)
                    pix = doc[n - 1].get_pixmap(dpi=OCR_DPI, colorspace=pymupdf.csGRAY)
                    imgs.append(Image.frombytes("L", (pix.width, pix.height), pix.samples))
                for n, text in zip(batch, pool.map(ocr, imgs)):
                    cache[str(n)] = text
                    done += 1
                    _report(0.02 + 0.68 * done / total, f"OCR: page {done} of {total}", "extract")
        cache_file.write_text(json.dumps(cache))
    return {n: _clean_ocr(cache[str(n)]) for n in page_numbers}


_LABEL_ADDR = re.compile(
    r"(?im)^[ \t]*(?:(?:registered|office|home|residential|mailing|postal|permanent|current|correspondence|billing|shipping|business|head[ ]?office|"
    r"company|employee|street)[ \t]+)?address(?:[ \t]+line[ \t]*\d)?[ \t]*[:=\u2013\u2014-]+[ \t]*(?P<a>[^\n]{6,160})")
_LABEL_PIN = re.compile(r"(?i)\b(?:pin(?:[ ]?code)?|zip(?:[ ]?code)?|postal[ ]?code|post[ ]?code)[ \t]*[:=\u2013\u2014-]*[ \t]*(?P<z>[A-Z0-9][A-Z0-9 -]{3,8}\d)\b")


def _phone_ok(s: str) -> bool:
    return 7 <= len(re.sub(r"\D", "", s)) <= 15


_LABEL_PHONE = re.compile(
    r"(?i)(?<![A-Za-z])(?P<lab>m|t|tel|telephone|ph|phone|mobile|mob|cell|fax|contact(?:[ \t]+no)?)(?![A-Za-z])[ \t]*[:.\-]*[ \t]*"
    r"(?P<v>\+?\(?\d[\d \t().-]{5,}\d)")


def _luhn_ok(num: str) -> bool:
    d = [int(c) for c in re.sub(r"\D", "", num)]
    if not 13 <= len(d) <= 19:
        return False
    s = 0
    for i, x in enumerate(reversed(d)):
        if i % 2:
            x = x * 2 - 9 if x * 2 > 9 else x * 2
        s += x
    return s % 10 == 0


_ADDR_TAIL = r"(?:,[ \t]*[A-Z0-9][^,\n]{1,38}){0,4}"      # ", Suite 400, New York, NY 10022, USA"
_STREET_SFX = ("street|st|road|rd|avenue|ave|boulevard|blvd|lane|ln|drive|dr|court|ct|place|pl|way|square|sq|terrace|ter|"
               "highway|hwy|parkway|pkwy|plaza|marg|nagar|colony")
_PH_ONLY = re.compile(r"\[[A-Z_]+_\d+(?:_R2)?\]")           # an already-redacted tag must never be detected again
_PATTERNS = [
    ("EMAIL", r"\b[\w.+-]+@[\w-]+(?:\.[\w-]+)+\b"),
    ("SSN", r"\b\d{3}-\d{2}-\d{4}\b"),
    ("AADHAAR", r"\b\d{4}\s\d{4}\s\d{4}\b"),
    ("PAN", r"\b[A-Z]{5}\d{4}[A-Z]\b"),
    ("CREDIT_CARD", r"\b(?:\d[ -]?){13,19}\b"),
    ("EMPLOYEE_ID", r"\b(?:EMP|EID)[-_ ]?\d{3,8}\b"),
    # postal addresses: "120 Park Avenue, Suite 400, New York, NY 10022", "Plot 45, Road No. 12, Hyderabad 500034", "PO Box 123"
    ("ADDRESS", r"(?<![\w-])\d{1,6}[A-Za-z]?(?:[-/]\d{1,5})?[ \t,]+(?:[A-Z0-9][\w.'-]*[ \t]+){1,4}?(?i:" + _STREET_SFX + r")\b\.?" + _ADDR_TAIL),
    ("ADDRESS", r"(?i:\b(?:plot|flat|door|house|h\.?[ ]?no\.?|block|building|bldg|floor|suite|ste|unit|apt|apartment|room)\b[ \t.#:-]*(?:no\.?[ \t#:-]*)?(?=[\w/-]*\d)[\w/-]+)" + _ADDR_TAIL),
    ("ADDRESS", r"(?i:\bp\.?[ ]?o\.?[ ]?box[ ]+\d+)" + _ADDR_TAIL),
    # phone numbers: international (+1 (555) 0174, +91 98765 43210, +44 20 7946 0958), US (555) 123-4567, 555-123-4567, Indian mobiles
    ("PHONE", r"(?<![\w+])\+\d{1,3}[ \t.-]*(?:\(\d{1,4}\)|\d{1,4})(?:[ \t.-]*\d{2,5}){1,4}(?![\w-])"),
    ("PHONE", r"(?<![\w-])\(\d{3}\)[ \t.-]*\d{3,4}(?:[ \t.-]*\d{4})?(?![\w-])"),
    ("PHONE", r"(?<![\w-])\d{3}[ .-]\d{3}[ .-]\d{4}(?![\w-])"),
    ("PHONE", r"(?<![\w-])0\d{2,4}[ -]\d{3,4}[ -]\d{3,4}(?![\w-])"),
    ("PHONE", r"(?<![\w.-])(?:(?:\+?91|0)[ -]?)?[6-9]\d{4}[ -]?\d{5}(?![\w.-])"),
]
_nlp = None

# Words that are NOT part of a person's name (headings, labels, business words). A name containing one is rejected.
_GENERIC = set("""
tier tiers workflow workflows metadata navigate navigation vendor vendors description descriptions key keys process processes policy policies
control controls risk risks data system systems report reports review reviews approval approvals level levels step steps section sections
table tables page pages slide slides summary overview introduction appendix version versions status type types name names date dates title
titles number numbers id ids project projects team teams member members leader leaders manager managers department departments company
document documents file files form forms template templates guide guides manual standard standards framework security cyber network
networks access user users admin administrator service services platform platforms application applications tool tools model models plan
plans strategy objective objectives scope purpose requirement requirements details detail notes note action actions owner owners approver
reviewer email phone address client customer customers employee employees staff student students course btech cse business continuity
management governance compliance assessment audit incident response recovery general information technology software hardware cloud
infrastructure operations operational financial finance legal regulatory enterprise global group limited ltd inc corp llc pvt private public
confidential internal external draft final yes no tbd na nil none null nan pending open closed high medium low critical priority category
categories item items list lists total amount count code codes event events log logs alert alerts rule rules ticket tickets queue
monday tuesday wednesday thursday friday saturday sunday january february march april may june july august september october november december
""".split())
_GENERIC |= set("""
the a an and or of to for with by in on at from as is are was were be been being this that these those it its our your their his her my we you
they he she them us who whom which what when where why how all any each every some not nor so than too very can will would should could may
might must shall do does did has have had also but if then else about into over under between among per via
recycle bin discard delete remove provider providers folder folders button buttons menu icon click select search home settings
everyone anyone someone others other another same above below following respective relevant designated assigned concerned responsible
authorised authorized applicable board committee council office tba tbc
""".split())
_PARTICLES = {"de", "van", "von", "bin", "al", "el", "la", "da", "di", "del", "dos", "bhai"}
_LABEL_STOP = {"enroll", "enrollment", "roll", "reg", "registration", "course", "dept", "department", "branch", "year", "sem",
               "semester", "email", "phone", "mobile", "contact", "id", "no", "number", "date", "on", "at", "of", "and", "is", "for"}
_TOKEN_RE = re.compile(r"[A-Za-z][A-Za-z.'\u2019-]*")
_NAME_VAL = r"[A-Za-z][A-Za-z.'\u2019-]*(?:[ \t]+[A-Za-z][A-Za-z.'\u2019-]*){0,3}"
_LABEL_NAME = re.compile(
    r"(?im)(?:^|[|\t,;(]|[ \t]{2,})[ \t]*(?:(?:full|employee|student|candidate|applicant|customer|client|member|leader|lead|mentor|"
    r"supervisor|manager|owner|author|reviewer|approver|assignee|guide|faculty|contact(?:[ \t]+person)?|team[ \t]+leader)"
    r"(?:[ \t]*\d+)?(?:[ \t]*name)?|name)[ \t]*[:=\u2013\u2014-]+[ \t]*(?P<n>" + _NAME_VAL + r")")
_BY_NAME = re.compile(
    r"(?i)\b(?:prepared|reviewed|approved|submitted|signed|reported|created|verified|authored|assigned|written|presented)"
    r"[ \t]+(?:by|to)[ \t:]+(?P<n>" + _NAME_VAL + r")")
_LABEL_ID = re.compile(
    r"(?i)\b(?P<lab>(?:enroll(?:ment)?|roll|registration|reg|student|employee|emp|staff|admission|usn|uid)[ .]*(?:no|number|id|code|#)?"
    r"|id[ ]*(?:no|number)?)[ \t]*[:=\u2013\u2014-]+[ \t]*(?P<v>[A-Za-z0-9][A-Za-z0-9/_-]{3,})")


def _tidy(value: str, start: int):
    """Trim spaces and trailing punctuation from a name span; keeps the dot of an initial ('John A.')."""
    value = re.split(r"(?<=[a-z]{2})[.!?][ \t]", value.split("\n")[0])[0]      # a name ends where the sentence ends
    lead = len(value) - len(value.lstrip())
    value, start = value.strip(), start + lead
    while value and value[-1] in ".,;:'\u2019\"":
        last = value.split()[-1].rstrip(".,;:'\u2019\"") if value.split() else ""
        if value[-1] == "." and len(last) == 1:
            break
        value = value[:-1]
    return value, start


STRICT_NAMES = True     # True: a name guessed by the language model is kept only with supporting evidence (see _name_evidence)
_FIRST_NAMES = set("""
james john robert michael william david richard joseph thomas charles christopher daniel matthew anthony mark donald steven paul andrew joshua
kenneth kevin brian george timothy ronald edward jason jeffrey ryan jacob gary nicholas eric jonathan stephen larry justin scott brandon benjamin
samuel gregory alexander patrick frank raymond jack dennis jerry tyler aaron jose adam nathan henry zachary douglas peter kyle noah ethan jeremy
walter christian keith roger terry austin sean gerald carl harold dylan arthur lawrence jordan jesse bryan billy bruce gabriel joe logan albert
alan eugene russell vincent philip bobby johnny bradley mary patricia jennifer linda elizabeth barbara susan jessica sarah karen lisa nancy betty
margaret sandra ashley kimberly emily donna michelle carol amanda dorothy melissa deborah stephanie rebecca sharon laura cynthia kathleen amy angela
shirley anna brenda pamela emma nicole helen samantha katherine christine debra rachel carolyn janet catherine maria heather diane olivia julie
joyce victoria ruth virginia lauren kelly christina joan evelyn judith andrea hannah megan cheryl jacqueline martha madison teresa gloria sara
janice marie julia grace judy theresa rose beverly denise marilyn amber danielle abigail brittany diana jane lori alexis kayla
priya anita anil amit rahul rohan vikram vijay ajay sanjay sunil suresh ramesh mahesh ravi raj rajesh rakesh deepak manoj naveen praveen pradeep
kiran arjun karthik krishna ganesh sai venkat srinivas lakshmi sita radha divya pooja neha sneha swati kavita sunita anjali ishita nisha meera
rekha shweta aditi sujith vanshika naomi marta daniela pedro luis carlos juan miguel jorge antonio manuel francisco javier fernando sergio
ricardo pablo diego andres marco luca giovanni mario paolo stefan hans klaus jan piotr tomasz pawel marek eva ewa agnieszka katarzyna ahmed
mohammed muhammad ali omar hassan hussein fatima aisha yusuf ibrahim wei li zhang wang chen yang liu ying hiro kenji yuki akira sakura olga ivan
dmitri sergei natalia elena
""".split())
_TITLE_RX = re.compile(r"(?i)\b(?:mr|mrs|ms|miss|dr|prof|shri|smt|sir|madam)\.?[ \t]*$")
_EVIDENCE_RX = re.compile(r"[\w.+-]+@[\w-]+\.[\w.]+|\+?\d[\d ()-]{7,}\d|\b(?:EMP|EID)[-_ ]?\d+")
_PIC_MARK = "[Text found in pictures]"
_OTHER_MARKS = ("[Speaker notes]", "[Text boxes]", "[Headers and footers]", "[Form fields and comments]")


def _in_picture(text: str, pos: int) -> bool:
    return text.rfind(_PIC_MARK, 0, pos) > max(text.rfind(m, 0, pos) for m in _OTHER_MARKS)


def _name_evidence(text: str, start: int, end: int, value: str, emails) -> bool:
    """Why we believe a name guessed by the language model is really a person:
    a title in front (Dr., Ms.), a middle initial, an email / phone / ID on the same or a neighbouring line,
    or the name appearing inside an email address of the document."""
    if _TITLE_RX.search(text[max(0, start - 12):start]):
        return True
    toks = value.split()
    if len(toks) >= 3 and any(re.fullmatch(r"[A-Z]\.?", t) for t in toks[1:-1]):
        return True
    ls = text.rfind("\n", 0, start) + 1
    prev = text.rfind("\n", 0, max(ls - 1, 0)) + 1
    le = text.find("\n", end)
    le = len(text) if le == -1 else le
    nxt = text.find("\n", le + 1)
    nxt = len(text) if nxt == -1 else nxt
    if _EVIDENCE_RX.search(text[prev:nxt]):
        return True
    parts = [t.lower().strip(".'") for t in toks if len(t) >= 4]
    return any(t in e for e in emails for t in parts)


def _is_person_name(text: str) -> bool:
    """spaCy guess is kept only if it looks like a real full name: 2-4 capitalised words, none of them a business/label word."""
    toks = text.replace("\n", " ").split()
    if not 2 <= len(toks) <= 4:
        return False
    for t in toks:
        low = t.lower().strip(".,'\u2019-")
        if low in _GENERIC or any(ch.isdigit() for ch in t):
            return False
        if low in _PARTICLES:
            if 0 < toks.index(t) < len(toks) - 1:      # 'bin', 'van', 'de' only between two name words
                continue
            return False
        if not re.fullmatch(r"[A-Z][A-Za-z.'\u2019-]*", t):
            return False
    return True


def _trim_name(value: str):
    """Cut a captured name at the first label/business word; returns (name, offset) or None."""
    keep = []
    for m in _TOKEN_RE.finditer(value):
        low = m.group().lower().strip(".'\u2019-")
        if low in _GENERIC or low in _LABEL_STOP:
            break
        keep.append(m)
    if not keep or not keep[0].group()[0].isupper():
        return None
    name = value[keep[0].start():keep[-1].end()]
    if len(name.replace(".", "").replace(" ", "")) < 2:
        return None
    return name, keep[0].start()


_COL_NAME = ("name", "employee", "owner", "manager", "lead", "leader", "reviewer", "approver", "author", "contact", "assignee",
             "member", "mentor", "supervisor", "signatory", "student", "candidate")
_COL_SKIP = ("project", "file", "team", "company", "vendor", "product", "course", "document", "department", "app", "system",
             "control", "policy", "role", "job", "title", "organi", "client", "supplier", "service", "tool")


def _table_name_findings(pg):
    """Table rows written as 'a | b | c': if a header cell says Name / Employee / Owner ..., the cells below it are people.
    Only full names (2+ words) are taken, so vendor / product columns are not touched."""
    text, out = pg["text"], []
    header, pos = None, 0
    for line in text.split("\n"):
        line_start = pos
        pos += len(line) + 1
        if " | " not in line:
            header = None
            continue
        cells, off = [], 0
        for c in line.split(" | "):
            cells.append((c, line_start + off))
            off += len(c) + 3
        if header is None:
            low = [c.strip().lower() for c, _ in cells]
            header = [i for i, h in enumerate(low) if any(k in h for k in _COL_NAME) and not any(s in h for s in _COL_SKIP)]
            continue
        for i in header:
            if i >= len(cells):
                continue
            cell, off = cells[i]
            lead = len(cell) - len(cell.lstrip())
            got = _trim_name(cell.strip())
            if not got:
                continue
            name, st = _tidy(got[0], off + lead + got[1])
            if len(name.split()) >= 2:
                out.append({"type": "PERSON", "value": name, "page": pg["page"], "start": st, "end": st + len(name)})
    return out


_VERB_NAME = re.compile(
    r"\b(?i:call|contact|email|ask|ping|reach|meet|thanks|thank|dear|hi|hello|attn|attention|cc|mr|mrs|ms|dr|prof)\b[ \t:.]*"
    r"(?P<n>[A-Z][a-z]+(?:[ \t]+[A-Z][a-z.']+){1,2})")


def _row_name_findings(pg, found):
    """Pipe-separated rows: if one cell of the row is already a person / email / phone, other cells that look like
    full names (2-3 capitalised words, no business words) are people too."""
    text, out, pos = pg["text"], [], 0
    mine = [f for f in found if f["page"] == pg["page"] and f["type"] in ("PERSON", "EMAIL", "PHONE")]
    for line in text.split("\n"):
        ls = pos
        pos += len(line) + 1
        if " | " not in line or not any(ls <= f["start"] < ls + len(line) for f in mine):
            continue
        off = 0
        for c in line.split(" | "):
            cs = ls + off
            off += len(c) + 3
            val, st = _tidy(c, cs)
            if val and 2 <= len(val.split()) <= 3 and _is_person_name(val) and not any(f["start"] <= st < f["end"] for f in mine):
                out.append({"type": "PERSON", "value": val, "page": pg["page"], "start": st, "end": st + len(val)})
    return out


def _context_findings(pg):
    """Names and IDs recognised from the label in front of them: 'Name :- G Sujith', 'Leader: Vanshika jadaun',
    'Prepared by ...', 'Enroll No :- 2024...'. The label itself is NOT redacted, only the value."""
    text, out = pg["text"], _table_name_findings(pg)
    for rx in (_LABEL_NAME, _BY_NAME):
        for m in rx.finditer(text):
            got = _trim_name(m.group("n"))
            if got:
                name, st = _tidy(got[0], m.start("n") + got[1])
                if name:
                    out.append({"type": "PERSON", "value": name, "page": pg["page"], "start": st, "end": st + len(name)})
    for m in _VERB_NAME.finditer(text):                 # 'Call Priya Sharma', 'Thanks Anita Rao'
        v, st = _tidy(m.group("n"), m.start("n"))
        if _is_person_name(v):
            out.append({"type": "PERSON", "value": v, "page": pg["page"], "start": st, "end": st + len(v)})
    for m in _LABEL_ADDR.finditer(text):                # 'Address: ...' (continues on the next line only after a trailing comma)
        end, last, extra = m.end("a"), m.group("a"), 0
        for ln in text[end:].split("\n")[1:]:
            if extra >= 2 or not last.rstrip().endswith(",") or not ln.strip() or ":" in ln[:30] or len(ln) > 90:
                break
            end += 1 + len(ln)
            last, extra = ln, extra + 1
        val = text[m.start("a"):end].rstrip(" \t.,;")
        out.append({"type": "ADDRESS", "value": val, "page": pg["page"], "start": m.start("a"), "end": m.start("a") + len(val)})
    for m in _LABEL_PIN.finditer(text):                 # 'PIN: 520001', 'ZIP 94105'
        z = m.group("z").strip()
        if any(c.isdigit() for c in z):
            out.append({"type": "ADDRESS", "value": z, "page": pg["page"], "start": m.start("z"), "end": m.start("z") + len(z)})
    for m in _LABEL_PHONE.finditer(text):               # 'M +1 (555) 0174', 'Tel: 020 7946 0958'
        v = m.group("v").rstrip(" \t.-")
        if not _phone_ok(v):
            continue
        if len(m.group("lab")) <= 1 and not v.startswith(("+", "(")):
            continue                                       # single letter label (M / T): only trust clear phone shapes
        out.append({"type": "PHONE", "value": v, "page": pg["page"], "start": m.start("v"), "end": m.start("v") + len(v)})
    for m in _LABEL_ID.finditer(text):
        v = m.group("v")
        if not any(c.isdigit() for c in v):
            continue
        lab = m.group("lab").lower()
        kind = "EMPLOYEE_ID" if lab.startswith(("emp", "staff")) else ("ID_NUMBER" if lab.startswith("id") else "STUDENT_ID")
        out.append({"type": kind, "value": v, "page": pg["page"], "start": m.start("v"), "end": m.end("v")})
    return out


_nlp_lock = threading.Lock()
_warm_started = False


def _get_nlp():
    global _nlp
    with _nlp_lock:
        if _nlp is None:
            import spacy
            _nlp = spacy.load("en_core_web_sm",
                              disable=["parser", "lemmatizer", "attribute_ruler", "tagger"])
            _nlp.max_length = 2_000_000
    return _nlp


def warmup():
    """Load the name model in the background at app start, so the first upload is fast."""
    global _warm_started
    if _warm_started:
        return
    _warm_started = True

    def _load():
        try:
            _get_nlp()
        except Exception:
            pass
    threading.Thread(target=_load, daemon=True).start()


_MAILBOX = {"info", "support", "admin", "sales", "contact", "hr", "noreply", "careers", "help", "billing", "team", "office", "service",
            "services", "security", "compliance", "legal", "finance", "marketing", "press", "hello", "mail", "webmaster", "postmaster"}

_GENERIC |= _MAILBOX


def _email_name_findings(pg, emails):
    """An email like priya.nair@ or m.szewczyk@ tells us a person's name: redact 'Priya Nair' / 'Marta Szewczyk' in the text too."""
    out, text = [], pg["text"]
    for e in set(emails):
        parts = [x for x in re.split(r"[._-]+", e.split("@")[0]) if x.isalpha()]
        if len(parts) < 2 or len(parts[-1]) < 3:
            continue
        first, last = parts[0].lower(), parts[-1].lower()
        if {first, last} & (_MAILBOX | _GENERIC):
            continue
        if len(first) == 1:
            rx = rf"\b[A-Za-z][a-z]+[ \t]+(?:[A-Z]\.?[ \t]+)?{re.escape(last)}\b|\b{first}\.?[ \t]+{re.escape(last)}\b"
        else:
            rx = rf"\b{re.escape(first)}[ \t]+(?:[A-Z]\.?[ \t]+)?{re.escape(last)}\b"
        for m in re.finditer(rx, text, re.I):
            name = m.group()
            if all(t[0].isupper() for t in name.split()):
                out.append({"type": "PERSON", "value": name, "page": pg["page"], "start": m.start(), "end": m.end()})
    return out


def _fallback_detect(pages, use_ner=True):
    found = []
    emails = [m.lower() for pg in pages for m in re.findall(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+", pg["text"])]
    for pg in pages:
        for typ, rx in _PATTERNS:
            for m in re.finditer(rx, pg["text"]):
                if typ == "CREDIT_CARD" and not _luhn_ok(m.group()):
                    continue
                if typ == "PHONE" and not _phone_ok(m.group()):
                    continue
                found.append({"type": typ, "value": m.group(), "page": pg["page"],
                              "start": m.start(), "end": m.end()})
        found += _context_findings(pg)
        found += _email_name_findings(pg, emails)
    try:  # names via spaCy, then filtered so only real full names survive
        if not use_ner:
            raise RuntimeError
        docs = _get_nlp().pipe([p["text"] for p in pages], batch_size=8)
        for pg, doc in zip(pages, docs):
            for ent in doc.ents:
                raw = ent.text.lstrip()                      # a name never spans two lines or two sentences
                raw = re.split(r"(?<=[a-z]{2})[.!?][ \t]", raw.split("\n")[0])[0]
                val, st = _tidy(raw, ent.start_char + (len(ent.text) - len(raw)))
                if not val:
                    continue
                if ent.label_ == "PERSON" and _is_person_name(val):
                    text = pg["text"]
                    keep = (not STRICT_NAMES) or _name_evidence(text, st, st + len(val), val, emails) \
                        or (not _in_picture(text, st) and val.split()[0].lower() in _FIRST_NAMES)
                    if keep:
                        found.append({"type": "PERSON", "value": val, "page": pg["page"], "start": st, "end": st + len(val)})
                elif ent.label_ == "ORG" and REDACT_ORGANIZATIONS and _is_person_name(val):
                    found.append({"type": "ORGANIZATION", "value": val, "page": pg["page"], "start": st, "end": st + len(val)})
    except Exception:
        pass
    for pg in pages:
        found += _row_name_findings(pg, found)
    return found


def _fallback_redact(pages, findings):
    counters, seen, mapping = {}, {}, {}
    for f in sorted(findings, key=lambda x: (x["page"], x["start"])):
        key = (f["type"], f["value"].lower())
        if key not in seen:
            counters[f["type"]] = counters.get(f["type"], 0) + 1
            ph = f"[{f['type']}_{counters[f['type']]}]"
            seen[key] = ph
            mapping[ph] = f["value"]
    out = []
    for pg in pages:
        text = pg["text"]
        mine = [x for x in findings if x["page"] == pg["page"]]
        for f in sorted(mine, key=lambda x: x["start"], reverse=True):
            text = text[:f["start"]] + seen[(f["type"], f["value"].lower())] + text[f["end"]:]
        out.append({"page": pg["page"], "text": text})
    return {"redacted_pages": out, "mapping": mapping}


# ------------------------------------------------------------ steps --------
def _call_extractor(fn, path: Path):
    kwargs = {}
    try:
        if "progress" in inspect.signature(fn).parameters:
            kwargs["progress"] = lambda done, total: _report(
                0.02 + 0.68 * done / max(total, 1), f"Reading page {done} of {total}", "extract")
    except (TypeError, ValueError):
        pass
    return fn(str(path), **kwargs)


def _extract(path: Path, used: dict):
    hook = EXTRACTOR_HOOKS.get(path.suffix.lower())
    fn = _hook(*hook) if hook else None
    used["Extraction"] = "teammate module" if fn else "temporary fallback"
    pages = _norm_pages(_call_extractor(fn, path) if fn else _fallback_extract(path))
    if fn and path.suffix.lower() == ".pdf" and not any(p["text"].strip() for p in pages):
        pages = _norm_pages(_fallback_extract(path))  # scanned PDF: OCR fallback
        used["Extraction"] = "teammate module + fallback OCR"
    return pages


def _detect(pages, used: dict | None = None, quick: bool = False):
    hooks = [(m, n) for m, n, fast in DETECTOR_HOOKS if fast or not quick]
    fns = [f for f in (_hook(m, n) for m, n in hooks) if f]
    if used is not None:
        used["Detection"] = "teammate module" if fns else "temporary fallback"
    raw = []
    for fn in fns:
        raw += _norm_findings(fn(pages))
    if not fns:
        raw = _fallback_detect(pages, use_ner=not quick)
    return _resolve_overlaps(raw)


def _verify(red_pages, known):
    """Quick safety check on the redacted text:
    (1) re-run the fast pattern detectors, (2) look for any value we already found that is still there."""
    found = _detect(red_pages, quick=True)
    types = {f["value"]: f["type"] for f in known if len(f["value"].strip()) >= 3}
    for pg in red_pages:
        for v, t in types.items():
            for m in re.finditer(re.escape(v), pg["text"]):
                found.append({"type": t, "value": v, "page": pg["page"], "start": m.start(), "end": m.end()})
    return _resolve_overlaps(found)


def _redact(pages, findings, used: dict | None = None):
    fn = _hook(*REDACTOR_HOOK)
    if used is not None:
        used["Redaction"] = "teammate module" if fn else "temporary fallback"
    res = fn(pages, findings) if fn else _fallback_redact(pages, findings)
    return _norm_pages(res["redacted_pages"]), dict(res.get("mapping", {}))


# ------------------------------------------------------------ main ---------
def run_pipeline(file_bytes: bytes, filename: str, progress=None) -> dict:
    global _progress_cb
    _progress_cb = progress
    file_type = detect_type(filename)
    if not file_type:
        raise ValueError("Unsupported file type. Use: " + ", ".join(SUPPORTED))
    if not file_bytes:
        raise ValueError("The file is empty.")

    fid = uuid.uuid4().hex[:8]
    orig_path = ORIG_DIR / f"{fid}_{Path(filename).name}"
    orig_path.write_bytes(file_bytes)

    used, timings = {}, {}
    t = time.time()
    _report(0.01, "Reading the file...", "extract")
    pages = _extract(orig_path, used)
    timings["Extraction"] = round(time.time() - t, 1)
    if not any(p["text"].strip() for p in pages):
        raise ValueError("No text found in this file.")

    t = time.time()
    _report(0.72, "Scanning for names, IDs, emails, phones...", "detect")
    findings = _detect(pages, used)
    timings["Detection"] = round(time.time() - t, 1)

    t = time.time()
    _report(0.85, "Replacing PII with safe tags...", "redact")
    red_pages, mapping = _redact(pages, findings, used)
    total = len(findings)
    timings["Redaction"] = round(time.time() - t, 1)

    # Verification: quick scan of the redacted text; redact anything still found.
    t = time.time()
    _report(0.92, "Scanning the redacted output again...", "verify")
    known = list(findings)
    remaining = _verify(red_pages, known)
    for _ in range(2):
        if not remaining:
            break
        total += len(remaining)
        known += remaining
        red_pages, extra = _redact(red_pages, remaining)
        for k, v in extra.items():
            mapping[k if k not in mapping else k + "_R2"] = v
        remaining = _verify(red_pages, known)
    timings["Verification"] = round(time.time() - t, 1)

    pct = 100.0 if total == 0 else round((total - len(remaining)) / total * 100, 1)
    red_text = "\n\n".join(f"--- Page/Slide {p['page']} ---\n{p['text']}" for p in red_pages)
    (RED_DIR / f"{fid}_redacted.txt").write_text(red_text, encoding="utf-8")
    _report(1.0, "Done", "done")

    return {
        "id": fid, "name": filename, "file_type": file_type, "orig_path": str(orig_path),
        "pages": pages, "redacted_pages": red_pages, "redacted_text": red_text,
        "findings": findings, "mapping": mapping, "total": total,
        "remaining": len(remaining), "redaction_pct": pct, "modules": used,
        "timings": timings, "page_count": len(pages),
    }